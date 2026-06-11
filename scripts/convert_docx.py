"""Convert a .docx file to Markdown.

Usage:
    uv run --with python-docx python scripts/convert_docx.py <in.docx> <out.md>
    # or: python scripts/convert_docx.py <in.docx> <out.md>
"""
import sys
from pathlib import Path


def docx_to_markdown(docx_path: str, md_path: str) -> None:
    from docx import Document  # type: ignore

    doc = Document(docx_path)
    lines: list[str] = []

    for para in doc.paragraphs:
        style = para.style.name if para.style else ""
        text = para.text.strip()

        if not text:
            lines.append("")
            continue

        if style.startswith("Heading 1"):
            lines.append(f"# {text}")
        elif style.startswith("Heading 2"):
            lines.append(f"## {text}")
        elif style.startswith("Heading 3"):
            lines.append(f"### {text}")
        elif style.startswith("Heading 4"):
            lines.append(f"#### {text}")
        elif style in ("List Bullet", "List Bullet 2", "List Paragraph"):
            lines.append(f"- {text}")
        elif style in ("List Number", "List Number 2"):
            lines.append(f"1. {text}")
        else:
            lines.append(text)

    for table in doc.tables:
        lines.append("")
        header = True
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if header:
                lines.append("| " + " | ".join(["---"] * len(cells)) + " |")
                header = False
        lines.append("")

    Path(md_path).write_text("\n".join(lines), encoding="utf-8")
    print(f"Written: {md_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(f"Usage: python {sys.argv[0]} <in.docx> <out.md>")
        sys.exit(1)
    docx_to_markdown(sys.argv[1], sys.argv[2])
